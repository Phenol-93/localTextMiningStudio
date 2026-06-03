"""Document import helpers."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from charset_normalizer import from_path
from docx import Document as DocxDocument

from app.db.database import connection_scope, initialize_database
from app.db.repositories import DocumentRepository


SUPPORTED_EXTENSIONS = {".txt", ".csv", ".docx"}


class DocumentImportError(Exception):
    """Raised when a document cannot be imported."""


@dataclass(frozen=True)
class ImportedDocument:
    """Parsed document payload before it is stored in SQLite."""

    filename: str
    file_path: Path
    file_type: str
    raw_text: str
    metadata: dict[str, Any]


def import_document(
    database_path: Path | str,
    project_id: int,
    file_path: Path | str,
    *,
    text_column: str | None = None,
) -> int:
    """Parse a supported document and insert it into the documents table."""
    parsed = parse_document(file_path, text_column=text_column)
    initialize_database(database_path)

    with connection_scope(database_path) as connection:
        repository = DocumentRepository(connection)
        return repository.create(
            project_id,
            parsed.filename,
            filename=parsed.filename,
            file_path=parsed.file_path,
            file_type=parsed.file_type,
            raw_text=parsed.raw_text,
            metadata=parsed.metadata,
            status="imported",
        )


def parse_document(file_path: Path | str, *, text_column: str | None = None) -> ImportedDocument:
    """Parse a supported local document without writing it to the database."""
    path = Path(file_path)
    if not path.exists():
        raise DocumentImportError(f"文件不存在：{path}")

    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentImportError(f"暂不支持的文件类型：{extension or '无扩展名'}")

    try:
        if extension == ".txt":
            raw_text, metadata = _parse_txt(path)
        elif extension == ".csv":
            raw_text, metadata = _parse_csv(path, text_column=text_column)
        else:
            raw_text, metadata = _parse_docx(path)
    except DocumentImportError:
        raise
    except Exception as error:
        raise DocumentImportError(f"解析失败：{error}") from error

    return ImportedDocument(
        filename=path.name,
        file_path=path.resolve(),
        file_type=extension.removeprefix("."),
        raw_text=raw_text,
        metadata=metadata,
    )


def get_csv_columns(file_path: Path | str) -> list[str]:
    """Return CSV column names for UI selection."""
    path = Path(file_path)
    try:
        sample = _read_text_file(path)
        reader = csv.reader(sample.splitlines())
        first_row = next(reader, [])
    except Exception as error:
        raise DocumentImportError(f"读取 CSV 列失败：{error}") from error

    return [column.strip() or f"column_{index + 1}" for index, column in enumerate(first_row)]


def _parse_txt(path: Path) -> tuple[str, dict[str, Any]]:
    raw_text = _read_text_file(path)
    return raw_text, {"source": "txt", "size_bytes": path.stat().st_size}


def _parse_csv(path: Path, *, text_column: str | None = None) -> tuple[str, dict[str, Any]]:
    content = _read_text_file(path)
    lines = content.splitlines()
    if not lines:
        return "", {"source": "csv", "columns": [], "text_column": None, "row_count": 0}

    reader = csv.DictReader(lines)
    columns = list(reader.fieldnames or [])
    rows = list(reader)
    if not columns:
        return "", {"source": "csv", "columns": [], "text_column": None, "row_count": 0}

    selected_column = text_column if text_column in columns else _select_default_text_column(columns, rows)
    raw_text = "\n".join(
        value
        for row in rows
        for value in [row.get(selected_column, "").strip()]
        if value
    )

    return raw_text, {
        "source": "csv",
        "columns": columns,
        "text_column": selected_column,
        "row_count": len(rows),
    }


def _parse_docx(path: Path) -> tuple[str, dict[str, Any]]:
    document = DocxDocument(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    raw_text = "\n".join(paragraphs)
    return raw_text, {"source": "docx", "paragraph_count": len(paragraphs)}


def _select_default_text_column(columns: list[str], rows: list[dict[str, str]]) -> str:
    for column in columns:
        if any(row.get(column, "").strip() for row in rows):
            return column
    return columns[0]


def _read_text_file(path: Path) -> str:
    match = from_path(path).best()
    if match is not None:
        return _normalize_newlines(str(match))

    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return _normalize_newlines(path.read_text(encoding=encoding))
        except UnicodeDecodeError:
            continue

    raise DocumentImportError("无法识别文本编码")


def _normalize_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")

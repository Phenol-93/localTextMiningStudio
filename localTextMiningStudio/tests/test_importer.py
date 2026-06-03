from docx import Document as DocxDocument

from app.core.importer import import_document
from app.db import database
from app.db.repositories import DocumentRepository, ProjectRepository


def _create_project(database_path, project_dir):
    database.initialize_database(database_path)
    with database.connection_scope(database_path) as connection:
        return ProjectRepository(connection).create("导入测试项目", project_dir)


def _get_document(database_path, document_id):
    with database.connection_scope(database_path) as connection:
        return DocumentRepository(connection).get(document_id)


def test_import_txt(tmp_path) -> None:
    database_path = tmp_path / "project.sqlite"
    project_id = _create_project(database_path, tmp_path)
    sample_file = tmp_path / "sample.txt"
    sample_file.write_text("第一行\n第二行", encoding="utf-8")

    document_id = import_document(database_path, project_id, sample_file)

    document = _get_document(database_path, document_id)
    assert document["filename"] == "sample.txt"
    assert document["file_type"] == "txt"
    assert document["raw_text"] == "第一行\n第二行"
    assert document["metadata_json"]


def test_import_csv_defaults_to_first_text_column(tmp_path) -> None:
    database_path = tmp_path / "project.sqlite"
    project_id = _create_project(database_path, tmp_path)
    sample_file = tmp_path / "sample.csv"
    sample_file.write_text("text,category\n苹果很好吃,food\n香蕉也很好吃,food\n", encoding="utf-8")

    document_id = import_document(database_path, project_id, sample_file)

    document = _get_document(database_path, document_id)
    assert document["filename"] == "sample.csv"
    assert document["file_type"] == "csv"
    assert document["raw_text"] == "苹果很好吃\n香蕉也很好吃"
    assert '"text_column": "text"' in document["metadata_json"]


def test_import_docx(tmp_path) -> None:
    database_path = tmp_path / "project.sqlite"
    project_id = _create_project(database_path, tmp_path)
    sample_file = tmp_path / "sample.docx"

    docx = DocxDocument()
    docx.add_paragraph("第一段")
    docx.add_paragraph("第二段")
    docx.save(sample_file)

    document_id = import_document(database_path, project_id, sample_file)

    document = _get_document(database_path, document_id)
    assert document["filename"] == "sample.docx"
    assert document["file_type"] == "docx"
    assert document["raw_text"] == "第一段\n第二段"
